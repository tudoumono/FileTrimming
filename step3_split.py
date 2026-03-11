"""
steps/step3_split.py - Phase 3: トークン推定 & 物理分割

ファイルサイズ（トークン数）を検査し、
OpenAI のコンテキスト上限を超えるファイルを物理的に分割する。
"""

import json
import shutil
import logging
from pathlib import Path

from .base import BaseStep
from utils.file_utils import detect_file_type, read_text_file
from utils.token_utils import estimate_tokens

logger = logging.getLogger(__name__)


def _extract_text_for_estimation(filepath: Path, ftype: str) -> str:
    """トークン推定のためにテキストを抽出する（簡易版）"""
    if ftype == "text":
        return read_text_file(filepath)

    if ftype == "excel":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append("\t".join(str(c) if c is not None else "" for c in row))
            wb.close()
            return "\n".join(parts)
        except Exception:
            return ""

    if ftype == "word":
        try:
            import docx
            doc = docx.Document(str(filepath))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return ""

    if ftype == "pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(filepath))
            return "\n".join(page.get_text() for page in doc)
        except Exception:
            return ""

    # その他: バイナリを無視して空文字
    try:
        return read_text_file(filepath)
    except Exception:
        return ""


def _split_text_by_lines(text: str, token_limit: int, overlap_lines: int = 20) -> list[str]:
    """テキストを行単位で分割（オーバーラップ付き）"""
    lines = text.split("\n")
    chunks = []
    start = 0

    while start < len(lines):
        end = start + 1
        current = "\n".join(lines[start:end])

        while end < len(lines) and estimate_tokens(current) < token_limit:
            end += 1
            current = "\n".join(lines[start:end])

        # 超過した場合は1行戻す
        if estimate_tokens(current) > token_limit and end > start + 1:
            end -= 1

        chunks.append("\n".join(lines[start:end]))
        start = max(start + 1, end - overlap_lines)

    return chunks


class Step3Split(BaseStep):
    step_number = 3
    step_name = "トークン推定 & 物理分割"

    def execute(self):
        prev_files_dir = self.config.paths.step_dir(2) / "files"
        if not prev_files_dir.exists():
            raise FileNotFoundError("Step 2 の出力が見つかりません。")

        out_dir = self.step_dir / "files"
        out_dir.mkdir(parents=True, exist_ok=True)

        token_limit = self.config.processing.token_limit
        split_log = []

        for src in sorted(prev_files_dir.rglob("*")):
            if not src.is_file() or src.name.startswith("."):
                continue

            rel = src.relative_to(prev_files_dir)
            ftype = detect_file_type(src)

            try:
                raw_text = _extract_text_for_estimation(src, ftype)
                tokens = estimate_tokens(raw_text) if raw_text else 0
            except Exception:
                tokens = 0

            entry = {
                "file": str(rel),
                "type": ftype,
                "estimated_tokens": tokens,
                "split": False,
                "parts": [],
            }

            if tokens <= token_limit or tokens == 0:
                # 分割不要 → そのままコピー
                dest = out_dir / rel
                dest.parent.mkdir(parents=True, exist_ok=True)

                if self.should_process_file(str(rel), dest):
                    shutil.copy2(src, dest)
                    self.manifest.mark_file_done(str(rel))

                entry["parts"] = [str(rel)]
            else:
                # 分割が必要
                logger.info("  分割対象: %s (%d tokens > %d limit)", rel, tokens, token_limit)
                entry["split"] = True

                chunks = _split_text_by_lines(raw_text, token_limit)
                stem = src.stem
                suffix = ".txt"  # 分割後はテキストとして保存

                for i, chunk in enumerate(chunks):
                    part_name = f"{stem}_part{i:03d}{suffix}"
                    part_rel = rel.parent / part_name
                    part_dest = out_dir / part_rel
                    part_dest.parent.mkdir(parents=True, exist_ok=True)

                    if self.should_process_file(str(part_rel), part_dest):
                        part_dest.write_text(chunk, encoding="utf-8")
                        self.manifest.mark_file_done(str(part_rel))

                    entry["parts"].append(str(part_rel))

                logger.info("    → %d パートに分割", len(chunks))

            split_log.append(entry)

        # 分割ログ出力
        (self.step_dir / "split_log.json").write_text(
            json.dumps(split_log, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        total = len(split_log)
        split_count = sum(1 for e in split_log if e["split"])
        logger.info("  合計: %d ファイル（うち %d 件を分割）", total, split_count)
